"""Pruebas de carga de archivos, consulta de indicadores, estudiantes y reporte Excel vía API."""

import io
import unittest

from openpyxl import load_workbook

from tests.helpers import admin_headers, client, csv_bytes, reset_database, upload

MAT_H = ["Tipo de documento", "Número de documento", "Programa"]
EGR_H = ["Número de identificación", "Programa", "Fecha y Hora de Grado"]


class IngestionTests(unittest.TestCase):
    def setUp(self):
        reset_database()
        self.c = client()
        self.h = admin_headers(self.c)

    def _mat(self, periodo, docs, prog="IS", tipo="matriculados", name=None):
        name = name or f"{'Matriculados Nuevos' if tipo == 'nuevos' else 'Matriculados'} {periodo}.csv"
        return upload(self.c, self.h, tipo, periodo, name, csv_bytes(MAT_H, [("CC", d, prog) for d in docs]))

    def test_full_flow(self):
        self.assertEqual(self._mat("2020-1", ["1001", "1002", "1003", "1004"]).status_code, 200)
        self.assertEqual(self._mat("2020-1", ["1001", "1002", "1003", "1004"], tipo="nuevos").status_code, 200)
        for p, docs in [("2020-2", ["1001", "1002", "1003", "1004"]), ("2021-1", ["1002", "1003", "1004"]),
                        ("2021-2", ["1003", "1004"]), ("2022-1", ["1004"])]:
            r = self._mat(p, docs)
            self.assertEqual(r.status_code, 200, r.text)
        # Egresados anual 2020 con fecha ISO en texto (antes se interpretaba con el día primero)
        r = upload(self.c, self.h, "egresados", "2020", "Egresados 2020.csv",
                   csv_bytes(EGR_H, [("1001", "Ing. De sistemas", "2020-12-10")]))
        self.assertEqual(r.status_code, 200, r.text)
        ind = self.c.get("/api/indicators?nivel=pregrado&vista=ies", headers=self.h).json()
        self.assertEqual(ind["periodos"], ["2020-1", "2020-2", "2021-1", "2021-2", "2022-1"])
        tda = {x["periodo"]: x for x in ind["tda"]}
        self.assertEqual(tda["2022-1"]["desertores"], 1)   # 1002 (1003 solo falta un periodo)
        tga = {x["semestre"]: x["graduados"] for x in ind["tga"]}
        self.assertEqual(tga[2], 1)                         # grado en diciembre = semestre 2 (S2)
        # Estudiantes
        s = self.c.get("/api/students/summary?nivel=pregrado", headers=self.h).json()
        self.assertEqual(s["resumen"]["Graduado"]["cantidad"], 1)
        self.assertEqual(s["resumen"]["Desertor"]["cantidad"], 1)
        d = self.c.get("/api/students/1001", headers=self.h).json()
        self.assertEqual(d["estado"], "Graduado")
        self.assertEqual(d["grados"][0]["periodo"], "2020-2")
        # Reporte Excel
        r = self.c.get("/api/indicators/export?nivel=pregrado&vista=ies", headers=self.h)
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Resumen", "TDA", "TAI", "TDC", "TDPA", "TGA"])
        self.assertGreaterEqual(len(wb["TDA"]._charts), 1)
        self.assertEqual(wb["TDA"]["A6"].value, "2021-1")
        self.assertAlmostEqual(wb["TDA"]["D6"].value, 0.0)
        self.assertAlmostEqual(wb["TDA"]["D8"].value, 1 / 3, places=4)

    def test_sequence_must_be_consecutive(self):
        self.assertEqual(self._mat("2020-1", ["1001"]).status_code, 200)
        r = self._mat("2021-1", ["1001"])
        self.assertEqual(r.status_code, 400)
        self.assertIn("consecutivos", r.json()["detail"])
        self.assertEqual(self._mat("2019-2", ["1001"]).status_code, 200)    # historia hacia atrás

    def test_nuevos_require_matriculados(self):
        r = self._mat("2020-1", ["1001"], tipo="nuevos")
        self.assertEqual(r.status_code, 400)

    def test_duplicate_requires_replace(self):
        self._mat("2020-1", ["1001", "1002"])
        self.assertEqual(self._mat("2020-1", ["1001"]).json()["error"]["code"], "DATA_002")
        r = upload(self.c, self.h, "matriculados", "2020-1", "Matriculados 2020-1.csv",
                   csv_bytes(MAT_H, [("CC", "1001", "IS")]), reemplazar=True)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["registros"], 1)

    def test_filename_period_mismatch(self):
        r = self._mat("2020-1", ["1001"], name="Matriculados 2-2020.csv")
        self.assertEqual(r.status_code, 400)

    def test_graduation_date_outside_semester(self):
        self._mat("2020-1", ["1001"])
        r = upload(self.c, self.h, "egresados", "2020-1", "Egresados 1-2020.csv",
                   csv_bytes(EGR_H, [("1001", "IS", "2020-08-15")]))
        self.assertEqual(r.status_code, 400)
        self.assertIn("2020-08", r.json()["detail"].replace("15/08/2020", "2020-08"))

    def test_annual_and_semester_graduates_cannot_overlap(self):
        self._mat("2020-1", ["1001"])
        upload(self.c, self.h, "egresados", "2020", "Egresados 2020.csv", csv_bytes(EGR_H, [("1001", "IS", "2020-03-01")]))
        r = upload(self.c, self.h, "egresados", "2020-1", "Egresados 1-2020.csv", csv_bytes(EGR_H, [("1001", "IS", "2020-03-01")]))
        self.assertEqual(r.json()["error"]["code"], "DATA_002")

    def test_invalid_documents_rejected(self):
        r = upload(self.c, self.h, "matriculados", "2020-1", "Matriculados 2020-1.csv",
                   csv_bytes(MAT_H, [("CC", "ABC-12$", "IS"), ("CC", "", "IS")]))
        self.assertEqual(r.status_code, 400)

    def test_unknown_program_confirmation_flow(self):
        content = csv_bytes(MAT_H, [("CC", "1001", "Ingeniería Espacial"), ("CC", "1002", "Ing. De sistemas"),
                                    ("CC", "1003", "Esp. En Gerencia Lunar")])
        r = upload(self.c, self.h, "matriculados", "2020-1", "Matriculados 2020-1.csv", content)
        body = r.json()
        self.assertEqual(body["status"], "confirmar_programas")
        sugeridos = {p["nombre"]: p["nivel_sugerido"] for p in body["programas"]}
        self.assertEqual(sugeridos, {"Ingeniería Espacial": "pregrado", "Esp. En Gerencia Lunar": "posgrado"})
        progs = {p["nombre"]: p["id"] for p in self.c.get("/api/programs", headers=self.h).json()}
        decisiones = [
            {"nombre": "Ingeniería Espacial", "accion": "crear", "nivel": "pregrado"},
            {"nombre": "Esp. En Gerencia Lunar", "accion": "asociar", "programa_id": progs["Especialización en Gerencia Financiera"]},
        ]
        r = upload(self.c, self.h, "matriculados", "2020-1", "Matriculados 2020-1.csv", content, decisiones=decisiones)
        self.assertEqual(r.json()["status"], "success", r.text)
        names = {p["nombre"]: p["nivel"] for p in self.c.get("/api/programs", headers=self.h).json()}
        self.assertEqual(names["Ingeniería Espacial"], "pregrado")
        pos = self.c.get("/api/indicators?nivel=posgrado", headers=self.h).json()
        self.assertEqual(pos["conteos"][0]["matriculados"], 1)

    def test_delete_middle_period_blocked(self):
        for p in ("2020-1", "2020-2", "2021-1"):
            self._mat(p, ["1001"])
        loads = {l["periodo"]: l["id"] for l in self.c.get("/api/loads", headers=self.h).json()}
        self.assertEqual(self.c.delete(f"/api/loads/{loads['2020-2']}", headers=self.h).status_code, 400)
        self.assertEqual(self.c.delete(f"/api/loads/{loads['2021-1']}", headers=self.h).status_code, 200)

    def test_status_and_warnings(self):
        self._mat("2020-1", ["1001"])
        st = self.c.get("/api/loads/status", headers=self.h).json()
        self.assertEqual(st["siguiente"], {"tipo": "nuevos", "periodo": "2020-1"})
        self.assertTrue(any("Nuevos" in w for w in st["advertencias"]))

    def test_level_change_updates_indicators(self):
        self._mat("2020-1", ["1001"], prog="IS")
        progs = {p["nombre"]: p["id"] for p in self.c.get("/api/programs", headers=self.h).json()}
        self.assertEqual(self.c.get("/api/indicators?nivel=pregrado", headers=self.h).json()["conteos"][0]["matriculados"], 1)
        self.c.put(f"/api/programs/{progs['Ingeniería de Sistemas']}", headers=self.h, json={"nivel": "posgrado"})
        self.assertEqual(self.c.get("/api/indicators?nivel=pregrado", headers=self.h).json()["conteos"], [])


if __name__ == "__main__":
    unittest.main()
