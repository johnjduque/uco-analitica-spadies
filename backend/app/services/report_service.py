"""
REPORTE EN EXCEL DE LOS INDICADORES

Genera un libro .xlsx con una hoja de resumen y una hoja por indicador (TDA,
TAI, TDC, TDPA, TGA), cada una con su tabla y una gráfica nativa de Excel.
Los porcentajes se escriben como números (formato %), listos para analizar.
"""

import io
from datetime import datetime
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GREEN = "047857"
GREEN_LIGHT = "ECFDF5"
GREY = "64748B"
_HEADER_FILL = PatternFill("solid", fgColor=GREEN)
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="CBD5E1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_PCT = "0.00%"

NIVEL_TXT = {"pregrado": "Pregrado", "posgrado": "Posgrado", "todos": "Pregrado y posgrado"}
VISTA_TXT = {"ies": "Institución (IES)", "programa": "Programa académico"}

DEFINICIONES = {
    "TDA": "Tasa de deserción anual: % de matriculados en t-2 que en t llevan dos periodos sin matrícula (t-1 y t) y no se han graduado.",
    "TAI": "Tasa de ausencia intersemestral: % de matriculados en el periodo t que no se matriculan en t+1 y no se han graduado.",
    "TDC": "Tasa de deserción por cohorte: % acumulado de primíparos de la cohorte que abandonaron (sin volver a matricularse ni graduarse), asignados al semestre de su última matrícula.",
    "TDPA": "Tasa de deserción promedio acumulada: desertores acumulados de todas las cohortes que alcanzan el semestre s, sobre sus primíparos.",
    "TGA": "Tasa de graduación acumulada: graduados acumulados de todas las cohortes que alcanzan el semestre s, sobre sus primíparos.",
}


def _frac(v: Optional[float]):
    return None if v is None else round(v / 100.0, 6)


def _title(ws, text: str, subtitle: str) -> None:
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=GREEN)
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color=GREY)
    ws["A2"].alignment = Alignment(wrap_text=False)


def _header(ws, row: int, headers: List[str]) -> None:
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill, cell.font, cell.border = _HEADER_FILL, _HEADER_FONT, _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _body_cell(ws, row, col, value, fmt=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.border = _BORDER
    if fmt:
        cell.number_format = fmt
    return cell


def _widths(ws, widths: Dict[int, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _line_chart(title: str, y_title: str, x_title: str, height=8.5, width=22) -> LineChart:
    ch = LineChart()
    ch.title = title
    ch.y_axis.title = y_title
    ch.x_axis.title = x_title
    ch.y_axis.number_format = "0%"
    ch.height, ch.width = height, width
    ch.legend.position = "b"
    return ch


def _period_sheet(wb, name: str, title: str, filters_txt: str, series: List[Dict], count_key: str, count_label: str):
    ws = wb.create_sheet(name)
    _title(ws, title, filters_txt)
    ws["A3"] = DEFINICIONES[name]
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A3:F3")
    ws.row_dimensions[3].height = 32
    _header(ws, 5, ["Periodo", "Base (matriculados)", count_label, "Tasa", "Retención"])
    r = 6
    for item in series:
        _body_cell(ws, r, 1, item["periodo"])
        _body_cell(ws, r, 2, item["base"], "#,##0")
        _body_cell(ws, r, 3, item.get(count_key), "#,##0")
        _body_cell(ws, r, 4, _frac(item["tasa"]), _PCT)
        _body_cell(ws, r, 5, _frac(item["retencion"]), _PCT)
        r += 1
    _widths(ws, {1: 12, 2: 20, 3: 16, 4: 12, 5: 12})
    if series:
        ch = _line_chart(f"{name} por periodo", "Tasa", "Periodo")
        ch.add_data(Reference(ws, min_col=4, min_row=5, max_row=r - 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=6, max_row=r - 1))
        ws.add_chart(ch, "H5")
    ws.freeze_panes = "B6"


def _semester_sheet(wb, name: str, title: str, filters_txt: str, series: List[Dict], count_key: str, count_label: str):
    ws = wb.create_sheet(name)
    _title(ws, title, filters_txt)
    ws["A3"] = DEFINICIONES[name]
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A3:F3")
    ws.row_dimensions[3].height = 32
    _header(ws, 5, ["Semestre", "Tasa", count_label, "Primíparos", "Cohortes"])
    r = 6
    for item in series:
        _body_cell(ws, r, 1, f"S{item['semestre']}")
        _body_cell(ws, r, 2, _frac(item["tasa"]), _PCT)
        _body_cell(ws, r, 3, item.get(count_key), "#,##0")
        _body_cell(ws, r, 4, item["primiparos"], "#,##0")
        _body_cell(ws, r, 5, item["cohortes"], "#,##0")
        r += 1
    _widths(ws, {1: 11, 2: 11, 3: 14, 4: 13, 5: 11})
    if series:
        ch = _line_chart(f"{name} por semestre", "Tasa", "Semestre")
        ch.add_data(Reference(ws, min_col=2, min_row=5, max_row=r - 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=6, max_row=r - 1))
        ws.add_chart(ch, "H5")
    return ws, r


def _cohort_matrix(ws, start_row: int, title: str, matrix: Dict, chart_title: str):
    rows = matrix["cohortes"]
    n_sem = max((len(c["valores"]) for c in rows), default=0)
    ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, color=GREEN)
    hdr = start_row + 1
    _header(ws, hdr, ["Cohorte", "Primíparos"] + [f"S{s}" for s in range(1, n_sem + 1)])
    r = hdr + 1
    for c in rows:
        _body_cell(ws, r, 1, c["cohorte"])
        _body_cell(ws, r, 2, c["primiparos"], "#,##0")
        for s, v in enumerate(c["valores"], start=3):
            _body_cell(ws, r, s, _frac(v), _PCT)
        r += 1
    if rows and n_sem:
        rng = f"C{hdr + 1}:{get_column_letter(2 + n_sem)}{r - 1}"
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                          end_type="max", end_color="F87171"))
        # Gráfica: hasta las 8 cohortes más recientes con datos (una línea por cohorte)
        ch = _line_chart(chart_title, "Tasa acumulada", "Semestre", height=9, width=24)
        recent = list(range(hdr + 1, r))[-8:]
        for row_i in recent:
            ch.add_data(Reference(ws, min_col=2 + 1, max_col=2 + n_sem, min_row=row_i), from_rows=True, titles_from_data=False)
        for serie, row_i in zip(ch.series, recent):
            serie.tx = SeriesLabel(v=str(ws.cell(row=row_i, column=1).value))
        ch.set_categories(Reference(ws, min_col=3, max_col=2 + n_sem, min_row=hdr))
        ws.add_chart(ch, f"A{r + 2}")
    for col in range(3, 3 + n_sem):
        ws.column_dimensions[get_column_letter(col)].width = 8.5
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 12
    return r


def build_report(result: Dict, institution: str, generated_by: str, program_label: str) -> bytes:
    f = result.get("filtros", {})
    filters_txt = (f"{institution} · Nivel: {NIVEL_TXT.get(f.get('nivel'), f.get('nivel'))} · "
                   f"Vista: {VISTA_TXT.get(f.get('vista'), f.get('vista'))} · Programas: {program_label}")
    wb = Workbook()

    # --- Resumen -------------------------------------------------------------
    ws = wb.active
    ws.title = "Resumen"
    _title(ws, "Indicadores de deserción y permanencia (metodología SPADIES 3.0)", filters_txt)
    meta = [
        ("Generado", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Generado por", generated_by),
        ("Periodos analizados", f"{result['periodos'][0]} a {result['periodos'][-1]} (ventana: últimos 10 años, máx. 20 semestres)" if result["periodos"] else "Sin datos"),
    ]
    for i, (k, v) in enumerate(meta, start=4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    _header(ws, 8, ["Periodo", "Matriculados", "Primíparos", "Graduados"])
    r = 9
    for c in result["conteos"]:
        _body_cell(ws, r, 1, c["periodo"])
        _body_cell(ws, r, 2, c["matriculados"], "#,##0")
        _body_cell(ws, r, 3, c["primiparos"], "#,##0")
        _body_cell(ws, r, 4, c["graduados"], "#,##0")
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Definiciones").font = Font(bold=True, color=GREEN)
    for k, v in DEFINICIONES.items():
        r += 1
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        ws.row_dimensions[r].height = 30
    if result.get("advertencias"):
        r += 2
        ws.cell(row=r, column=1, value="Advertencias de datos").font = Font(bold=True, color="B45309")
        for w in result["advertencias"]:
            r += 1
            ws.cell(row=r, column=1, value=f"• {w}")
    _widths(ws, {1: 20, 2: 16, 3: 14, 4: 14})

    # --- Hojas por indicador ----------------------------------------------------
    _period_sheet(wb, "TDA", "Tasa de Deserción Anual (TDA)", filters_txt, result["tda"], "desertores", "Desertores")
    _period_sheet(wb, "TAI", "Tasa de Ausencia Intersemestral (TAI)", filters_txt, result["tai"], "ausentes", "Ausentes")

    ws = wb.create_sheet("TDC")
    _title(ws, "Tasa de Deserción por Cohorte (TDC)", filters_txt)
    ws["A3"] = DEFINICIONES["TDC"]
    ws.merge_cells("A3:L3")
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[3].height = 32
    _cohort_matrix(ws, 5, "Deserción acumulada por cohorte", result["tdc"], "TDC por cohorte (cohortes recientes)")
    ws.freeze_panes = "C7"

    _semester_sheet(wb, "TDPA", "Tasa de Deserción Promedio Acumulada (TDPA)", filters_txt, result["tdpa"], "desertores", "Desertores")
    ws, end = _semester_sheet(wb, "TGA", "Tasa de Graduación Acumulada (TGA)", filters_txt, result["tga"], "graduados", "Graduados")
    _cohort_matrix(ws, end + 2, "Graduación acumulada por cohorte", result["tga_cohortes"], "TGA por cohorte (cohortes recientes)")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
